import io
import zipfile

from docx import Document
from docxtpl import DocxTemplate
from jinja2.exceptions import TemplateSyntaxError
from mailmerge import MailMerge
from rest_framework import exceptions

from . import models
from .jinja import get_jinja_env


class _MagicPlaceholder:
    def __init__(self, parent=None, name=None):
        self._parent = parent
        self._name = name
        self._reports = parent._reports if parent else set()

        if str(self):
            self._reports.add(str(self))

    @property
    def reports(self):
        return list(self._reports)

    def __getitem__(self, idx):
        if type(idx) is int:
            if idx >= 2:
                raise IndexError()
            return _MagicPlaceholder(parent=self, name=f"{self}[]")
        return _MagicPlaceholder(parent=self, name=f"{self}.{idx}".strip("."))

    def __getattr__(self, attr):
        return _MagicPlaceholder(parent=self, name=f"{self}.{attr}".strip("."))

    def __len__(self):
        return 2

    def __str__(self):
        return self._name if self._name else ""


class DocxValidator:
    def _validate_is_docx(self):
        try:
            Document(self.template)
        except (ValueError, zipfile.BadZipfile):
            raise exceptions.ParseError("not a valid docx file")
        finally:
            self.template.seek(0)

    def validate_template_syntax(self, available_placeholders=None):  # pragma: no cover
        raise NotImplementedError(
            "validate_template_syntax must be implemented in engine class"
        )

    def validate(self, available_placeholders=None, sample_data=None):
        self._validate_is_docx()
        self.validate_template_syntax(available_placeholders, sample_data)

    def validate_available_placeholders(
        self, used_placeholders, available_placeholders
    ):
        # We don't validate available_placeholders if it's not given
        if not available_placeholders:
            return

        available_placeholders = self._normalize_available_placeholders(
            available_placeholders
        )

        referenced_unavailable = "; ".join(
            sorted(set(used_placeholders) - set(available_placeholders))
        )
        if referenced_unavailable:
            raise exceptions.ValidationError(
                f"Template uses unavailable placeholders: {referenced_unavailable}"
            )

    def _normalize_available_placeholders(self, placeholders):
        available_placeholders = set(placeholders)
        # add all prefixes of placeholders, so users don't
        # have to add "foo" if they have "foo.bar" in the list
        for ph in placeholders:
            prefix = ""
            for word in ph.split("."):
                prefix = f"{prefix}.{word}" if prefix else word
                if prefix.endswith("[]"):
                    available_placeholders.add(prefix[:-2])
                available_placeholders.add(prefix)
        return available_placeholders


class DocxTemplateEngine(DocxValidator):
    def __init__(self, template):
        self.template = template

    def validate_template_syntax(self, available_placeholders=None, sample_data=None):
        try:
            doc = DocxTemplate(self.template)
            root = _MagicPlaceholder()
            env = get_jinja_env()
            ph = {
                name: root[name] for name in doc.get_undeclared_template_variables(env)
            }
            doc.render(ph, env)

            if sample_data:
                doc.render(sample_data, env)

            self.validate_available_placeholders(
                used_placeholders=root.reports,
                available_placeholders=available_placeholders,
            )

        except TemplateSyntaxError as exc:
            arg_str = ";".join(exc.args)
            raise exceptions.ValidationError(f"Syntax error in template: {arg_str}")

        finally:
            self.template.seek(0)

    def merge(self, data, buf):
        doc = DocxTemplate(self.template)

        doc.render(data, get_jinja_env())
        doc.save(buf)
        return buf


class DocxMailmergeEngine(DocxValidator):
    def __init__(self, template):
        self.template = template

    def _get_placeholders(self, document):
        return [
            field.attrib.get("name")
            for part in document.parts.values()
            for field in part.findall("//MergeField")
        ]

    def validate_template_syntax(self, available_placeholders=None, sample_data=None):
        document = MailMerge(self.template)
        # syntax can't be invalid as it's validated by office
        # suites. However we need to have *some* placeholders
        self.template.seek(0)
        used_placeholders = self._get_placeholders(document)
        self.validate_available_placeholders(
            used_placeholders=used_placeholders,
            available_placeholders=available_placeholders,
        )

        if sample_data:
            buffer = io.BytesIO()
            self.merge(sample_data, buffer)
            self.template.seek(0)

        if not len(used_placeholders):
            raise exceptions.ValidationError("Template has no merge fields")

    def merge(self, data, buf):
        with MailMerge(self.template) as document:
            document.merge(**data)
            document.write(buf)
            return buf


ENGINES = {
    models.Template.DOCX_TEMPLATE: DocxTemplateEngine,
    models.Template.DOCX_MAILMERGE: DocxMailmergeEngine,
}


def get_engine(engine, template):
    return ENGINES[engine](template)
